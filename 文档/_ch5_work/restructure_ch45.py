# -*- coding: utf-8 -*-
"""
按《研究报告old.pdf》第四、五章的风格与流程，重构研究报告.docx 的第4、5章。

old 第4章的二级流程：设计目标与总体思路 → 软件总体架构与数据流设计 →
计算实现 → 功能模块设计 → 研究流程设计 → 软件验证 → 电脑要求及配置。
其中「功能模块设计」（old 4.4）把各界面作为**三级标题**挂在父节下 —— 这正是
研究报告.docx 第5章要学的写法。

本脚本做四件事：
  1. 第4章：二级/三级标题改为真正的 Heading 2/3 样式，理顺层级与编号（4.1~4.6）；
  2. 第5章：新增「5.2 功能模块设计及其对耦合研究的支撑」父节，首页/教学/仿真/
     模型/优化/解释 降为其下的 5.2.1~5.2.6 三级标题；
  3. 修正第4章失效的内文图号引用（图49/图50、图4-11 等）与图注编号；
  4. 修正段落样式误用（带图的段落用「图片」、题注用「图注」）。

插图本身（w:drawing）不改，随所在段落整体移动；第5章的图号顺序不变（仍是 5-1~5-8）。
"""
import sys
from copy import deepcopy
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("研究报告.docx")
DST = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("研究报告.restructured.docx")

d = docx.Document(str(SRC))
body = d.element.body
P = d.paragraphs
OWNER = P[0]._parent


# --------------------------------------------------------------- 工具
def set_heading(i, style_name, text):
    """套用真正的 Heading 样式，并清掉与样式重复的直接格式。保留段内书签。"""
    p_el = P[i]._p
    docx.text.paragraph.Paragraph(p_el, OWNER).style = d.styles[style_name]
    pPr = p_el.get_or_add_pPr()
    for tag in ("w:spacing", "w:ind", "w:outlineLvl", "w:jc", "w:keepNext", "w:keepLines"):
        for e in pPr.findall(qn(tag)):
            pPr.remove(e)
    for r in p_el.findall(qn("w:r")):
        p_el.remove(r)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p_el.append(r)


def set_body_text(p_el, text):
    """整段改写正文，按文档约定拆分中英文 run（中文宋体 / 西文 Times New Roman）。"""
    import re
    CJK = re.compile(r"[⺀-鿿　-〿＀-￯]")

    def rpr(attrs):
        e = OxmlElement("w:rPr")
        rf = OxmlElement("w:rFonts")
        for k, v in attrs.items():
            rf.set(qn("w:" + k), v)
        e.append(rf)
        c = OxmlElement("w:color")
        c.set(qn("w:val"), "000000")
        e.append(c)
        return e

    for r in p_el.findall(qn("w:r")):
        p_el.remove(r)
    segs = []
    for ch in text:
        lat = not CJK.match(ch)
        if segs and segs[-1][1] == lat:
            segs[-1][0] += ch
        else:
            segs.append([ch, lat])
    for s, lat in segs:
        r = OxmlElement("w:r")
        r.append(rpr({"cs": "Times New Roman"} if lat else
                     {"ascii": "宋体", "hAnsi": "宋体", "cs": "Times New Roman"}))
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = s
        r.append(t)
        p_el.append(r)


def replace_in(i, old, new, required=True):
    p = P[i]
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new, 1)
            return
    full = "".join(r.text for r in p.runs)
    if old in full:
        set_body_text(p._p, full.replace(old, new, 1))
        return
    if required:
        raise SystemExit(f"[{i}] 未找到 {old!r}: {full[:50]!r}")


def clone(new_idx, text, heading=None):
    """从模板段落复制出新的独立段落（不带走模板书签）。"""
    el = deepcopy(P[new_idx]._p)
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for e in el.findall(qn(tag)):
            el.remove(e)
    for r in el.findall(qn("w:r")):
        el.remove(r)
    if heading:
        el.append(OxmlElement("w:r"))
        para = docx.text.paragraph.Paragraph(el, OWNER)
        para.style = d.styles[heading]
        pPr = el.get_or_add_pPr()
        for tag in ("w:spacing", "w:ind", "w:outlineLvl", "w:jc", "w:keepNext", "w:keepLines"):
            for e in pPr.findall(qn(tag)):
                pPr.remove(e)
        for r in el.findall(qn("w:r")):
            el.remove(r)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        r.append(t)
        el.append(r)
    else:
        set_body_text(el, text)
    return el


# --------------------------------------------------------------- 0. 前置校验
region = [P[i]._p for i in range(208, 300)]
assert len(region) == 92, len(region)
# 208..299 在 body 中必须是连续的纯段落，中间不能夹书签/表格
kids = list(body)
assert kids[kids.index(P[208]._p) - 1] is P[207]._p
for a, b in zip(region, region[1:]):
    assert kids.index(b) == kids.index(a) + 1, "第4~5章区间内出现非段落元素"
print("region check ok: 92 consecutive paragraphs")

# --------------------------------------------------------------- 1. 第4章
print("== 第4章 ==")

# 1.1 标题层级化（复用原有标题段，只改样式与文字）
set_heading(208, "Heading 2", "4.1 设计目标与总体思路")   # 原 4.1 设计思路(Normal)
set_heading(218, "Heading 2", "4.3 计算实现与程序开发")   # 原 4.3 程序代码编写
set_heading(214, "Heading 3", "4.3.1 程序开发流程")       # 原 4.2 程序开发流程
set_heading(229, "Heading 3", "4.3.3 器材建模")           # 原 4.3 器材建模(重号)
set_heading(225, "Heading 3", "4.2.2 平台总体架构")       # 原 4.1.2 平台架构
set_heading(233, "Heading 2", "4.4 UI 界面设计")          # 原 4.4 UI 界面
set_heading(244, "Heading 2", "4.5 程序封装与部署")       # 原 4.5 程序封装

# 1.2 图号重排：平台架构类图提前到 4.2，4-2~4-5 相应位移
replace_in(217, "图4-2", "图4-4")     # 程序开发流程
replace_in(221, "图4-3", "图4-5")     # 核心计算脚本
replace_in(224, "图4-4", "图4-2")     # 平台总体结构与模块分层
replace_in(228, "图4-5", "图4-3")     # 平台总体架构
replace_in(219, "图4-3", "图4-5")     # 内文引用
replace_in(222, "图4-4", "图4-2")     # 内文引用

# 1.3 失效引用与题注修正
replace_in(250, "图4-11", "图4-10")           # 打包输出文件实为 图4-10
replace_in(253, "图49", "图4-11")             # 旧编号残留（图49/图50 已不存在）
replace_in(253, "图50", "图4-12")
replace_in(247, "图4-9软件封装流程", "图4-9 软件封装流程")
replace_in(252, "图4-11独立安装程序界面", "图4-11 独立安装程序界面")
P[231].style = d.styles["图片"]                # 该段放的是图
P[252].style = d.styles["图注"]                # 该段是题注

# 1.4 新增段落
h42 = clone(214, "4.2 软件总体架构与数据流设计", "Heading 2")
h421 = clone(236, "4.2.1 模块分层与数据流", "Heading 3")
h423 = clone(236, "4.2.3 共享状态与可追溯机制", "Heading 3")
b423 = clone(226,
             "平台的工程快照、任务状态、数据集、活动模型和计算结果以共享对象的形式在模块之间传递。"
             "教学模块保存的场景可以写入当前工程；仿真模块产出的正式结果与工程快照绑定，"
             "记录其参数版本与计算精度；模型模块生成的数据集登记采样方式与样本清单；"
             "优化模块的候选方案写回工程后，由平台自动发起一次正式仿真复核。"
             "任一环节产生的参数、结果和状态都可回溯到对应的计算任务，便于重复计算和实验前检查。")
h432 = clone(236, "4.3.2 核心计算脚本", "Heading 3")
h46 = clone(214, "4.6 电脑要求及配置", "Heading 2")
b46 = clone(226,
            "平台以 Windows 桌面程序的形式交付，最低配置要求为：处理器为 Intel 或 AMD x86-64 处理器；"
            "操作系统为 Windows 11 及以上（64 位）；内存至少 8 GB，推荐 16 GB 及以上；"
            "硬盘空间 7 GB 及以上；显示分辨率建议 1920×1080 及以上，"
            "以便完整显示五栏工作流与结果工作区。以源码方式运行时，需要 Python 3.11～3.13 环境"
            "并安装 requirements/requirements.txt 中的依赖；使用安装包时不需要预先安装 Python，"
            "安装程序已将解释器与全部依赖库一并封装。")

ch4_order = [
    208,                                                        # 4.1
    209, 210, 211, 212, 213,
    h42,                                                        # 4.2
    h421, 222, 223, 224,                                        # 4.2.1
    225, 226, 227, 228,                                         # 4.2.2
    h423, b423,                                                 # 4.2.3
    218,                                                        # 4.3
    214, 215, 216, 217,                                         # 4.3.1
    h432, 219, 220, 221,                                        # 4.3.2
    229, 230, 231, 232,                                         # 4.3.3
    233, 234, 235,                                              # 4.4
    236, 237, 238, 239,                                         # 4.4.1
    240, 241, 242, 243,                                         # 4.4.2
    244, *range(245, 256),                                      # 4.5
    h46, b46,                                                   # 4.6
    256,                                                        # 章末空行
]

# --------------------------------------------------------------- 2. 第5章
print("== 第5章 ==")

set_heading(258, "Heading 2", "5.1 平台总体功能与工作流")   # 原 5.1 首页
set_heading(263, "Heading 3", "5.2.2 教学模块")
set_heading(268, "Heading 3", "5.2.3 仿真模块")
set_heading(277, "Heading 3", "5.2.4 模型模块")
set_heading(282, "Heading 3", "5.2.5 优化模块")
set_heading(287, "Heading 3", "5.2.6 解释模块")
set_heading(295, "Heading 2", "5.3 模块联动与 AI 辅助")     # 原 5.7

h51lead = clone(259,
                "平台按照“教学理解—正式仿真—参数研究—智能优化—结果解释”的闭环研究流程组织功能。"
                "首页按一级功能栏的顺序给出五个功能区入口，各模块围绕同一工程快照协同工作，"
                "使研究者不需要在页面之间反复切换即可完成从建模到结论的完整链路。")
h52 = clone(214, "5.2 功能模块设计及其对耦合研究的支撑", "Heading 2")
h521 = clone(236, "5.2.1 首页", "Heading 3")

ch5_order = [
    257,                                                        # Heading 1 保持章首
    258, h51lead, 260,                                          # 5.1
    h52,                                                        # 5.2
    h521, 259, 261, 262,                                        # 5.2.1 首页
    263, 264, 265, 266, 267,                                    # 5.2.2 教学模块
    268, *range(269, 277),                                      # 5.2.3 仿真模块
    277, *range(278, 282),                                      # 5.2.4 模型模块
    282, *range(283, 287),                                      # 5.2.5 优化模块
    287, *range(288, 295),                                      # 5.2.6 解释模块
    295, 296, 297, 298, 299,                                    # 5.3
]

# --------------------------------------------------------------- 3. 重排
used = [x for x in ch4_order + ch5_order if isinstance(x, int)]
assert sorted(used) == list(range(208, 300)), (
    "原文段落未完整覆盖", set(range(208, 300)) - set(used), [x for x in used if used.count(x) > 1])
assert len(used) == len(set(used)), "有段落被重复使用"

order_els = [P[x]._p if isinstance(x, int) else x for x in ch4_order + ch5_order]

for el in region:
    body.remove(el)
cursor = P[207]._p
for el in order_els:
    cursor.addnext(el)
    cursor = el

print(f"reordered {len(order_els)} paragraphs ({len(order_els) - len(used)} new)")

# --------------------------------------------------------------- 4. 图片段落样式归一
# 纯图段落统一用「图片」样式（居中），与用户已整理的段落保持一致。
fixed = 0
for para in d.paragraphs:
    if para._p in (P[207]._p,):
        continue
    if para._p in region and not para.text.strip() \
            and para._p.findall(".//" + qn("w:drawing")) \
            and para.style.name != "图片":
        para.style = d.styles["图片"]
        fixed += 1
print(f"normalized {fixed} image paragraphs to 图片 style")

d.save(str(DST))
print("saved ->", DST)
